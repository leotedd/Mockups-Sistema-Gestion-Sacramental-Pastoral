# -*- coding: utf-8 -*-
"""
Genera la presentación de documentación del MÓDULO CELEBRACIONES.

Fuente: mockup + las 18 evidencias de  evidencias/celebraciones/
Salida: documentacion/Modulo_Celebraciones_Mockup.pptx  (16:9, 18 diapositivas)

No modifica el mockup, ni las evidencias, ni el repositorio oficial.
Solo lee las imágenes PNG y construye el .pptx.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# ----------------------------------------------------------------------------
# Rutas
# ----------------------------------------------------------------------------
BUILD_DIR = os.path.dirname(os.path.abspath(__file__))
DOC_DIR = os.path.dirname(BUILD_DIR)
ROOT = os.path.dirname(DOC_DIR)
EV = os.path.join(ROOT, "evidencias", "celebraciones")
OUT = os.path.join(DOC_DIR, "Modulo_Celebraciones_Mockup.pptx")

def ev(name):
    p = os.path.join(EV, name)
    if not os.path.isfile(p):
        raise SystemExit("No se encontró la evidencia: " + p)
    return p

# ----------------------------------------------------------------------------
# Identidad visual (misma paleta del mockup)
# ----------------------------------------------------------------------------
BORGONA   = RGBColor(0x5C, 0x06, 0x1B)
BORGONA_D = RGBColor(0x4A, 0x05, 0x16)
BORGONA_L = RGBColor(0x7D, 0x1A, 0x30)
DORADO    = RGBColor(0xD4, 0xAF, 0x37)
DORADO_S  = RGBColor(0xF6, 0xEC, 0xCF)
DORADO_B  = RGBColor(0xC9, 0xA9, 0x4A)
MARFIL    = RGBColor(0xFA, 0xF7, 0xF0)
MARFIL2   = RGBColor(0xF1, 0xEC, 0xDF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
TXT       = RGBColor(0x1E, 0x1C, 0x1A)
TXT2      = RGBColor(0x43, 0x3E, 0x39)
TXT3      = RGBColor(0x6F, 0x6A, 0x64)
LINEA     = RGBColor(0xC3, 0xBB, 0xA8)
OKV       = RGBColor(0x2F, 0x6B, 0x3A)
PELIGRO   = RGBColor(0x9E, 0x22, 0x2B)

FONT = "Segoe UI"

SW, SH = 13.333, 7.5
MX = 0.62                     # margen lateral
CW = SW - 2 * MX              # ancho de contenido
BAND_H = 0.92                 # alto de la banda de título
CONTENT_TOP = BAND_H + 0.30
CONTENT_BOT = 7.06
FOOT_Y = 7.12

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]

# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def rect(slide, x, y, w, h, fill, line=None, line_w=1.0, rounded=False):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            sp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP):
    """paragraphs: lista de dicts {runs:[(t,size,bold,color,italic?)], align, bullet, level, sa, sb}"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", PP_ALIGN.LEFT)
        p.level = para.get("level", 0)
        if "sa" in para:
            p.space_after = Pt(para["sa"])
        if "sb" in para:
            p.space_before = Pt(para["sb"])
        if "lh" in para:
            p.line_spacing = para["lh"]
        runs = list(para["runs"])
        if para.get("bullet"):
            runs = [("•  ", runs[0][1], False, para.get("bulletcolor", DORADO_B))] + runs
        for rn in runs:
            t, size, bold, color = rn[0], rn[1], rn[2], rn[3]
            italic = rn[4] if len(rn) > 4 else False
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
    return tb


def slide_base(title, number, kicker="MÓDULO CELEBRACIONES"):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, SW, SH, MARFIL)                      # fondo
    rect(s, 0, 0, SW, BAND_H, BORGONA)                 # banda superior
    rect(s, 0, BAND_H, SW, 0.055, DORADO)              # filo dorado
    textbox(s, MX, 0.04, CW - 3.2, BAND_H - 0.04,
            [{"runs": [(title, 23, True, WHITE)]}], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        textbox(s, SW - MX - 3.2, 0.04, 3.2, BAND_H - 0.04,
                [{"runs": [(kicker, 10, True, DORADO_S)], "align": PP_ALIGN.RIGHT}],
                anchor=MSO_ANCHOR.MIDDLE)
    # pie
    rect(s, 0, FOOT_Y, SW, SH - FOOT_Y, MARFIL2)
    rect(s, 0, FOOT_Y, SW, 0.02, LINEA)
    textbox(s, MX, FOOT_Y, CW - 1.5, SH - FOOT_Y,
            [{"runs": [("Módulo Celebraciones · Sistema Integral de Gestión Sacramental y Pastoral",
                        8.5, False, TXT3)]}], anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, SW - MX - 1.5, FOOT_Y, 1.5, SH - FOOT_Y,
            [{"runs": [("%02d / 18" % number, 8.5, True, BORGONA_L)], "align": PP_ALIGN.RIGHT}],
            anchor=MSO_ANCHOR.MIDDLE)
    return s


def place_image(slide, path, bx, by, bw, bh, mat=True):
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    if bw / bh > ar:
        h = bh
        w = bh * ar
    else:
        w = bw
        h = bw / ar
    x = bx + (bw - w) / 2
    y = by + (bh - h) / 2
    if mat:
        rect(slide, x - 0.06, y - 0.06, w + 0.12, h + 0.12, WHITE, line=LINEA, line_w=1.0)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return (x, y, w, h)


def caption(slide, x, y, w, title, desc=None, align=PP_ALIGN.CENTER):
    paras = [{"runs": [(title, 12, True, BORGONA)], "align": align, "sa": 2}]
    if desc:
        paras.append({"runs": [(desc, 10.5, False, TXT2)], "align": align})
    textbox(slide, x, y, w, 0.9, paras)


def bullets(slide, x, y, w, h, items, size=13.5, gap=7, anchor=MSO_ANCHOR.TOP):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            text, lvl = it
        else:
            text, lvl = it, 0
        if lvl == 0:
            paras.append({"runs": [(text, size, False, TXT2)], "bullet": True, "sa": gap, "lh": 1.08})
        else:
            paras.append({"runs": [("– ", size - 1, False, DORADO_B), (text, size - 1, False, TXT2)],
                          "level": 1, "sa": max(2, gap - 3), "lh": 1.06})
    textbox(slide, x, y, w, h, paras, anchor=anchor)


def lead(slide, x, y, w, text, size=13.5):
    textbox(slide, x, y, w, 1.0,
            [{"runs": [(text, size, False, TXT2)], "lh": 1.12}])


# ============================================================================
# DIAPOSITIVA 1 — MÓDULO CELEBRACIONES
# ============================================================================
s = slide_base("Módulo Celebraciones", 1, kicker="")
textbox(s, MX, 1.55, CW, 1.0,
        [{"runs": [("Celebraciones parroquiales", 32, True, BORGONA)]}])
textbox(s, MX, 2.62, CW, 0.5,
        [{"runs": [("Programación sacramental y pastoral · Parroquia Santa Cruz, Chiquimulilla",
                    13, False, TXT3)]}])
rect(s, MX, 3.35, CW, 0.028, DORADO)

textbox(s, MX, 3.7, CW, 1.5, [{"runs": [(
    "El módulo Celebraciones permite administrar la programación de las celebraciones "
    "parroquiales, centralizando la información de fechas, horarios, lugares, celebrantes, "
    "intenciones y estados de cada celebración.", 15, False, TXT2)], "lh": 1.22}])

box = rect(s, MX, 5.15, CW, 1.55, MARFIL2, rounded=True)
rect(s, MX, 5.15, 0.09, 1.55, DORADO)
textbox(s, MX + 0.35, 5.28, CW - 0.7, 1.3, [
    {"runs": [("OBJETIVO GENERAL", 11, True, BORGONA)], "sa": 5},
    {"runs": [("Ofrecer a la secretaría parroquial una herramienta única para registrar, "
               "consultar, editar y dar seguimiento a todas las celebraciones, con una vista "
               "de listado y una vista de calendario, garantizando datos completos y coherentes.",
               13, False, TXT2)], "lh": 1.18},
])

# ============================================================================
# DIAPOSITIVA 2 — FLUJO GENERAL DEL MÓDULO
# ============================================================================
s = slide_base("Flujo general del módulo", 2)

def dbox(slide, x, y, w, h, txt, fill=MARFIL2, fg=BORGONA, size=11, bold=True, line=DORADO_B):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sp.adjustments[0] = 0.16
    except Exception:
        pass
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line; sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.05))
    for m in ("margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.02))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = txt
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = fg
    return sp

def arrow(slide, x, y, kind="r", w=0.34, h=0.26):
    shp = {"r": MSO_SHAPE.RIGHT_ARROW, "d": MSO_SHAPE.DOWN_ARROW, "l": MSO_SHAPE.LEFT_ARROW}[kind]
    sp = slide.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = DORADO
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

def seclabel(slide, x, y, txt):
    textbox(slide, x, y, 6.0, 0.34, [{"runs": [(txt, 11, True, BORGONA_L)]}])

bw, bh, gap = 2.13, 0.80, 0.40
aw, ah = 0.28, 0.24
xs = MX

def har(x, ycenter, kind="r"):
    arrow(s, x, ycenter - ah / 2, kind, aw, ah)

# ---- Flujo de registro (serpentina) ----
seclabel(s, MX, 1.34, "FLUJO DE REGISTRO")
y1 = 1.72
row1 = ["Listado de\ncelebraciones", "Nueva\ncelebración", "Datos\ngenerales", "Celebrantes", "Intenciones"]
for i, lb in enumerate(row1):
    x = xs + i * (bw + gap)
    dbox(s, x, y1, bw, bh, lb)
    if i < len(row1) - 1:
        har(x + bw + (gap - aw) / 2, y1 + bh / 2, "r")

xI = xs + 4 * (bw + gap)                         # columna de "Intenciones"
arrow(s, xI + bw / 2 - 0.12, y1 + bh + 0.05, "d", 0.24, 0.34)

y2 = y1 + bh + 0.60
dbox(s, xI, y2, bw, bh, "Guardar")
har(xI - gap + (gap - aw) / 2, y2 + bh / 2, "l")
xDet = xs + 3 * (bw + gap)
dbox(s, xDet, y2, bw, bh, "Detalle de\ncelebración", fill=DORADO_S, line=DORADO_B)
har(xDet - gap + (gap - aw) / 2, y2 + bh / 2, "l")

xEI = xs + 2 * (bw + gap)                        # Editar / Imprimir apilados
hh = 0.56
dbox(s, xEI, y2 + bh / 2 - hh - 0.04, bw, hh, "Editar")
dbox(s, xEI, y2 + bh / 2 + 0.04, bw, hh, "Imprimir")

# ---- separador ----
ysep = y2 + bh + 0.50
rect(s, MX, ysep, CW, 0.02, LINEA)

# ---- Flujo de consulta ----
seclabel(s, MX, ysep + 0.20, "FLUJO DE CONSULTA")
y3 = ysep + 0.62
consulta = [("Listado de\ncelebraciones", bw, MARFIL2),
            ("Calendario", bw, MARFIL2),
            ("Día · Laboral ·\nSemana · Mes", bw + 1.5, MARFIL2),
            ("Detalle de\ncelebración", bw + 0.15, DORADO_S)]
cx = xs
for i, (lb, w, fill) in enumerate(consulta):
    dbox(s, cx, y3, w, bh, lb, fill=fill,
         line=(DORADO_B if fill == DORADO_S else DORADO_B))
    if i < len(consulta) - 1:
        har(cx + w + (gap - aw) / 2, y3 + bh / 2, "r")
    cx += w + gap

# ============================================================================
# DIAPOSITIVA 3 — LISTADO DE CELEBRACIONES
# ============================================================================
s = slide_base("Listado de celebraciones", 3)
bullets(s, MX, CONTENT_TOP + 0.05, 4.15, 5.6, [
    "Consulta general de todas las celebraciones registradas.",
    "Búsqueda por texto: tipo, lugar, celebrante o intención.",
    "Filtros por rango de fechas (Desde / Hasta).",
    "Filtro por tipo de celebración y por lugar.",
    "Botones Buscar y Limpiar filtros.",
    "Columna Estado con distinción visual por color.",
    "Paginación de resultados.",
    "Acciones por fila: Ver y Editar.",
    "Acceso directo a Nueva celebración.",
], size=12.5, gap=6)
place_image(s, ev("01_listado_celebraciones.png"), 5.05, CONTENT_TOP, 7.66, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 4 — REGISTRO DE UNA CELEBRACIÓN
# ============================================================================
s = slide_base("Registro de una celebración", 4)
bullets(s, MX, CONTENT_TOP + 0.05, 4.55, 5.7, [
    "Formulario dividido en secciones: Datos generales, Celebrantes e Intenciones.",
    "Datos generales:",
    ("Tipo de celebración (obligatorio)", 1),
    ("Fecha (obligatorio)", 1),
    ("Hora desde (obligatorio)", 1),
    ("Hora hasta (opcional)", 1),
    ("Estado de la celebración", 1),
    ("Lugar (obligatorio)", 1),
    ("Observaciones (opcional)", 1),
    "Los campos obligatorios se señalan con asterisco (*).",
    "Acciones: Guardar, Guardar y nuevo, Cancelar.",
], size=12.5, gap=5)
place_image(s, ev("02_nueva_celebracion.png"), 5.35, CONTENT_TOP, 7.36, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 5 — GESTIÓN DE CELEBRANTES
# ============================================================================
s = slide_base("Gestión de celebrantes", 5)
lead(s, MX, CONTENT_TOP, CW,
     "Desde la sección Celebrantes se busca y selecciona a la persona que oficiará la "
     "celebración. Se puede registrar uno o varios celebrantes e identificar al principal.")
img_top = CONTENT_TOP + 0.95
img_h = CONTENT_BOT - img_top - 0.55
place_image(s, ev("03_seleccionar_celebrante.png"), MX, img_top, CW / 2 - 0.25, img_h)
place_image(s, ev("04_celebrante_agregado.png"), MX + CW / 2 + 0.25, img_top, CW / 2 - 0.25, img_h)
caption(s, MX, CONTENT_BOT - 0.5, CW / 2 - 0.25,
        "Seleccionar celebrante", "Buscador · Nombre · Tipo/rol · selector")
caption(s, MX + CW / 2 + 0.25, CONTENT_BOT - 0.5, CW / 2 - 0.25,
        "Celebrante agregado", "Principal (estrella) · Tipo/rol · acción Quitar")

# ============================================================================
# DIAPOSITIVA 6 — GESTIÓN DE INTENCIONES
# ============================================================================
s = slide_base("Gestión de intenciones", 6)
lead(s, MX, CONTENT_TOP, CW,
     "En la sección Intenciones se registran las intenciones solicitadas para la celebración. "
     "Una celebración puede contener todas las intenciones que se le registren.")
place_image(s, ev("05_agregar_intencion.png"), MX, img_top, CW / 2 - 0.25, img_h)
place_image(s, ev("06_intencion_agregada.png"), MX + CW / 2 + 0.25, img_top, CW / 2 - 0.25, img_h)
caption(s, MX, CONTENT_BOT - 0.5, CW / 2 - 0.25,
        "Agregar intención", "Tipo · Descripción · Solicitante · Observaciones")
caption(s, MX + CW / 2 + 0.25, CONTENT_BOT - 0.5, CW / 2 - 0.25,
        "Intención agregada", "Tipo · Descripción · Solicitante · acción Quitar")

# ============================================================================
# DIAPOSITIVA 7 — DETALLE DE CELEBRACIÓN
# ============================================================================
s = slide_base("Detalle de celebración", 7)
bullets(s, MX, CONTENT_TOP + 0.05, 4.3, 5.7, [
    "Vista de solo lectura con toda la información de la celebración.",
    "Información general: tipo, estado, fecha, horario y lugar.",
    "Observaciones registradas.",
    "Sacramentos asociados, cuando existen.",
    "Listado de celebrantes con su función (principal o concelebrante).",
    "Listado de intenciones con tipo, descripción y solicitante.",
    "Acciones: Volver, Imprimir, Eliminar y Editar.",
], size=12.5, gap=7)
place_image(s, ev("07_detalle_celebracion.png"), 5.2, CONTENT_TOP, 7.51, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 8 — EDICIÓN DE CELEBRACIÓN
# ============================================================================
s = slide_base("Edición de celebración", 8)
bullets(s, MX, CONTENT_TOP + 0.05, 4.55, 5.7, [
    "Mismo formulario que Nueva celebración, con los datos existentes ya cargados.",
    "Permite modificar los datos generales de la celebración.",
    "Gestión de celebrantes: agregar, quitar y cambiar el principal.",
    "Gestión de intenciones: agregar y quitar.",
    "Indicador de cambios sin guardar.",
    "Acciones: Cancelar y Guardar cambios.",
], size=12.5, gap=8)
place_image(s, ev("08_editar_celebracion.png"), 5.35, CONTENT_TOP, 7.36, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 9 — CALENDARIO DE CELEBRACIONES
# ============================================================================
s = slide_base("Calendario de celebraciones", 9)
bullets(s, MX, CONTENT_TOP + 0.05, 4.15, 5.6, [
    "Consulta visual de la programación parroquial.",
    "Cada celebración aparece con su hora y tipo.",
    "Al hacer clic sobre una celebración se abre su detalle.",
    "Vistas disponibles: Día, Laboral, Semana y Mes.",
    "Navegación: Anterior, Siguiente y Hoy.",
    "El día actual y las celebraciones canceladas se distinguen visualmente.",
], size=12.5, gap=8)
place_image(s, ev("09_calendario_mes.png"), 5.05, CONTENT_TOP, 7.66, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 10 — VISTAS DEL CALENDARIO
# ============================================================================
s = slide_base("Vistas del calendario", 10)
lead(s, MX, CONTENT_TOP, CW,
     "El calendario ofrece cuatro escalas de consulta. La vista Mes se muestra en la "
     "diapositiva anterior; a continuación se comparan las vistas Semana, Día y Laboral.")
gap10 = 0.34
col_w = (CW - 2 * gap10) / 3
data10 = [
    ("10_calendario_semana.png", "SEMANA", "Visualización de la programación semanal."),
    ("11_calendario_dia.png", "DÍA", "Programación detallada de una fecha."),
    ("12_calendario_laboral.png", "LABORAL", "Programación de los días laborables (lunes a viernes)."),
]
row_top = 2.95
for i, (img, t, d) in enumerate(data10):
    x = MX + i * (col_w + gap10)
    ix, iy, iw, ih = place_image(s, ev(img), x, row_top, col_w, col_w / 1.6)
    caption(s, x, iy + ih + 0.16, col_w, t, d)

# ============================================================================
# DIAPOSITIVA 11 — ESTADOS DE LAS CELEBRACIONES
# ============================================================================
s = slide_base("Estados de las celebraciones", 11)
lead(s, MX, CONTENT_TOP, 4.3,
     "El estado indica la situación actual de cada celebración dentro de la programación "
     "parroquial.")
estados = [
    ("Programada", "Registrada, aún sin confirmar.", RGBColor(0x1F, 0x5F, 0x8B)),
    ("Confirmada", "Verificada y lista para realizarse.", OKV),
    ("Realizada", "Ya se llevó a cabo.", TXT3),
    ("Cancelada", "No se realizará; se muestra tachada.", PELIGRO),
]
ey = CONTENT_TOP + 1.35
for nombre, desc, color in estados:
    rect(s, MX, ey, 0.16, 0.66, color)
    textbox(s, MX + 0.32, ey - 0.03, 3.9, 0.75, [
        {"runs": [(nombre, 13, True, color)], "sa": 1},
        {"runs": [(desc, 10.5, False, TXT2)]},
    ])
    ey += 0.86
place_image(s, ev("13_estado_cancelada.png"), 5.1, CONTENT_TOP, 7.61, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 12 — CONFIRMACIONES Y SEGURIDAD DE OPERACIONES
# ============================================================================
s = slide_base("Confirmaciones y seguridad de operaciones", 12)
bullets(s, MX, CONTENT_TOP + 0.05, 4.3, 5.4, [
    "Las operaciones que eliminan información requieren una confirmación previa.",
    "Antes de eliminar se muestra un resumen de la celebración afectada.",
    "El botón Eliminar recibe un tratamiento visual de acción peligrosa.",
    "La acción puede cancelarse sin efectos sobre los datos.",
    "Objetivo: reducir eliminaciones accidentales.",
], size=12.5, gap=9)
place_image(s, ev("14_confirmacion_eliminar.png"), 5.2, CONTENT_TOP, 7.51, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 13 — VALIDACIONES DEL FORMULARIO
# ============================================================================
s = slide_base("Validaciones del formulario", 13)
bullets(s, MX, CONTENT_TOP + 0.05, 4.3, 5.6, [
    "La validación se ejecuta al intentar guardar.",
    "Campos obligatorios: tipo de celebración, fecha, hora desde y lugar.",
    "La hora final debe ser posterior a la hora de inicio.",
    "Cada campo con error se resalta y muestra su mensaje.",
    "Se muestra una notificación general: «Revise el formulario».",
    "No se guarda la celebración hasta corregir los errores.",
], size=12.5, gap=8)
place_image(s, ev("15_validacion_formulario.png"), 5.1, CONTENT_TOP, 7.61, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 14 — RETROALIMENTACIÓN DEL SISTEMA
# ============================================================================
s = slide_base("Retroalimentación del sistema", 14)
lead(s, MX, CONTENT_TOP, CW,
     "El sistema informa el resultado de cada acción del usuario mediante notificaciones "
     "y estados de pantalla.")
place_image(s, ev("16_mensaje_exito.png"), MX, img_top, CW / 2 - 0.25, img_h)
place_image(s, ev("17_sin_resultados.png"), MX + CW / 2 + 0.25, img_top, CW / 2 - 0.25, img_h)
caption(s, MX, CONTENT_BOT - 0.5, CW / 2 - 0.25,
        "Operación exitosa", "Confirma cuando una acción se completa correctamente.")
caption(s, MX + CW / 2 + 0.25, CONTENT_BOT - 0.5, CW / 2 - 0.25,
        "Sin resultados", "Avisa cuando la búsqueda no encuentra registros.")

# ============================================================================
# DIAPOSITIVA 15 — IMPRESIÓN DE CELEBRACIÓN
# ============================================================================
s = slide_base("Impresión de celebración", 15)
bullets(s, MX, CONTENT_TOP + 0.05, 4.3, 5.5, [
    "Vista con la información preparada para impresión.",
    "Encabezado con los datos de la parroquia y el rango de fechas.",
    "Tabla con fecha, hora, celebración, lugar y celebrante.",
    "Presentación ordenada y formal de la programación.",
    "Acciones: Imprimir y Volver (no envía nada a una impresora física).",
], size=12.5, gap=9)
place_image(s, ev("18_impresion_celebracion.png"), 5.3, CONTENT_TOP, 7.41, CONTENT_BOT - CONTENT_TOP)

# ============================================================================
# DIAPOSITIVA 16 — REGLAS Y VALIDACIONES PRINCIPALES
# ============================================================================
s = slide_base("Reglas y validaciones principales", 16)
lead(s, MX, CONTENT_TOP, CW,
     "Reglas comprobadas en el mockup del módulo Celebraciones.")
reglas = [
    "Campos obligatorios antes de guardar.",
    "Hora final posterior a la hora de inicio.",
    "Selección obligatoria del tipo de celebración.",
    "Selección obligatoria del lugar.",
    "Confirmación previa antes de eliminar.",
    "Identificación de un único celebrante principal.",
    "Manejo de estados de la celebración.",
    "Mensajes de éxito al completar una operación.",
    "Estado «sin resultados» en búsquedas sin coincidencias.",
]
gx, gy = MX, CONTENT_TOP + 0.80
gw = (CW - 0.4) / 2
ghh = 0.80
for i, txt in enumerate(reglas):
    col = i % 2
    rowi = i // 2
    x = gx + col * (gw + 0.4)
    y = gy + rowi * (ghh + 0.14)
    rect(s, x, y, gw, ghh, MARFIL2, line=LINEA, line_w=0.75, rounded=True)
    rect(s, x + 0.18, y + ghh / 2 - 0.10, 0.20, 0.20, DORADO)
    textbox(s, x + 0.52, y, gw - 0.72, ghh,
            [{"runs": [(txt, 12, False, TXT2)], "lh": 1.05}], anchor=MSO_ANCHOR.MIDDLE)

# ============================================================================
# DIAPOSITIVA 17 — NAVEGACIÓN DEL MÓDULO
# ============================================================================
s = slide_base("Navegación del módulo", 17)

def nbox(x, y, w, h, txt, fill=MARFIL2, fg=BORGONA, size=11.5, line=DORADO_B, bold=True):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sp.adjustments[0] = 0.18
    except Exception:
        pass
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line; sp.line.width = Pt(1.2)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = txt
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = fg
    return sp

def connect(x1, y1, x2, y2):
    ln = s.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))  # 2 = elbow
    ln.line.color.rgb = DORADO_B
    ln.line.width = Pt(1.25)
    ln.shadow.inherit = False
    return ln

bwn, bhn = 2.5, 0.60
# raíz
rx = MX + 0.2
nbox(rx, 3.55, bwn, 0.7, "CELEBRACIONES", fill=BORGONA, fg=WHITE, line=BORGONA_D, size=12.5)
# nivel 2
l2x = rx + bwn + 1.0
y_lst, y_cal = 2.05, 4.95
nbox(l2x, y_lst, bwn, bhn, "Listado")
nbox(l2x, y_cal, bwn, bhn, "Calendario")
connect(rx + bwn, 3.90, l2x, y_lst + bhn / 2)
connect(rx + bwn, 3.90, l2x, y_cal + bhn / 2)
# nivel 3 - listado
l3x = l2x + bwn + 0.9
sub_list = ["Nueva celebración", "Ver detalle", "Filtros / búsqueda"]
ys_l = [1.38, 2.20, 3.02]
for t, yy in zip(sub_list, ys_l):
    nbox(l3x, yy, bwn, bhn, t, size=10.5)
    connect(l2x + bwn, y_lst + bhn / 2, l3x, yy + bhn / 2)
# nivel 4 - ver detalle
l4x = l3x + bwn + 0.85
sub_det = ["Editar", "Imprimir", "Eliminar"]
ys_d = [1.74, 2.42, 3.10]
for t, yy in zip(sub_det, ys_d):
    nbox(l4x, yy, bwn - 0.55, 0.54, t, size=10.5, fill=MARFIL)
    connect(l3x + bwn, ys_l[1] + bhn / 2, l4x, yy + 0.27)
# nivel 3 - calendario
sub_cal = ["Día", "Laboral", "Semana", "Mes"]
ys_c = [3.95, 4.62, 5.29, 5.96]
for t, yy in zip(sub_cal, ys_c):
    nbox(l3x, yy, bwn - 0.7, 0.54, t, size=10.5)
    connect(l2x + bwn, y_cal + bhn / 2, l3x, yy + 0.27)

# ============================================================================
# DIAPOSITIVA 18 — RESULTADO DEL DISEÑO
# ============================================================================
s = slide_base("Resultado del diseño", 18)
lead(s, MX, CONTENT_TOP, CW,
     "El mockup del módulo Celebraciones deja definidos los siguientes elementos de diseño:")
elems = ["Estructura visual", "Navegación entre pantallas", "Formularios y secciones",
         "Tablas y columnas", "Acciones y botones", "Modales",
         "Validaciones", "Calendario y sus vistas", "Estados de celebración",
         "Vista de impresión"]
cx, cy = MX, CONTENT_TOP + 0.95
cw = (CW - 2 * 0.3) / 3
chh = 0.6
for i, t in enumerate(elems):
    col = i % 3
    rowi = i // 3
    x = cx + col * (cw + 0.3)
    y = cy + rowi * (chh + 0.18)
    rect(s, x, y, cw, chh, MARFIL2, line=LINEA, line_w=0.75, rounded=True)
    textbox(s, x + 0.2, y, cw - 0.3, chh,
            [{"runs": [(t, 11.5, True, BORGONA)], "lh": 1.0}], anchor=MSO_ANCHOR.MIDDLE)

box = rect(s, MX, 5.95, CW, 1.02, MARFIL2, rounded=True)
rect(s, MX, 5.95, 0.09, 1.02, DORADO)
textbox(s, MX + 0.35, 5.95, CW - 0.7, 1.02, [{"runs": [(
    "El diseño funcional del módulo Celebraciones queda establecido como referencia para su "
    "posterior implementación e integración en el frontend del Sistema Integral de Gestión "
    "Sacramental y Pastoral.", 12.5, False, TXT2)], "lh": 1.15}], anchor=MSO_ANCHOR.MIDDLE)

# ----------------------------------------------------------------------------
prs.save(OUT)
print("OK ->", OUT)
print("Diapositivas:", len(prs.slides))
