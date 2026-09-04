# -*- coding: utf-8 -*-
"""Verificación estructural del PPTX generado."""
import os
from io import BytesIO
from pptx import Presentation
from PIL import Image

BUILD = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.dirname(BUILD)
PPTX = os.path.join(DOC, "Modulo_Celebraciones_Mockup.pptx")
EMU = 914400

p = Presentation(PPTX)
SW, SH = p.slide_width, p.slide_height
print("Archivo:", PPTX)
print("Diapositivas:", len(p.slides))
print("Tamano:", round(SW / EMU, 3), "x", round(SH / EMU, 3), "in (16:9 =",
      round(SW / SH, 4), ")")

issues = []
pics = 0
tol = 12000  # ~0.013 in de tolerancia de borde

for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
            continue
        L, T, R, B = sh.left, sh.top, sh.left + sh.width, sh.top + sh.height
        if L < -tol or T < -tol or R > SW + tol or B > SH + tol:
            issues.append("slide %d: elemento fuera de limites  L=%.2f T=%.2f R=%.2f B=%.2f  (%s)"
                          % (i, L / EMU, T / EMU, R / EMU, B / EMU, sh.shape_type))
        if sh.shape_type == 13:  # PICTURE
            pics += 1
            try:
                w, h = Image.open(BytesIO(sh.image.blob)).size
                src_ar = w / h
                box_ar = sh.width / sh.height
                if abs(src_ar - box_ar) / src_ar > 0.02:
                    issues.append("slide %d: imagen deformada  origen=%.3f  caja=%.3f"
                                  % (i, src_ar, box_ar))
            except Exception as e:
                issues.append("slide %d: no se pudo leer imagen (%s)" % (i, e))

print("Imagenes insertadas:", pics)
if issues:
    print("\nPROBLEMAS DETECTADOS:")
    for x in issues:
        print("  -", x)
    raise SystemExit(1)
print("\nOK: ninguna imagen deformada, ningun elemento fuera de la diapositiva.")
